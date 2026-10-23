"""Style the Pandoc Word reference for the portable EKZNW data_dev template.

Run with python-docx installed. This maintains styles, not report content.
"""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


path = Path(sys.argv[1])
doc = Document(path)
olive, brown, cream = "495C3F", "554838", "E2DFD0"
for style in doc.styles:
    if style.type == 1:
        style.font.name = "Arial"
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        style.font.color.rgb = RGBColor.from_string("000000")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.widow_control = True
for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
    style = doc.styles[name]
    style.paragraph_format.keep_with_next = False
    style.paragraph_format.space_after = Pt(6 if name != "Compact" else 3)
for level in range(1, 10):
    style = doc.styles[f"Heading {level}"]
    style.font.size = Pt(15 if level == 1 else 12 if level == 2 else 11)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(olive if level == 1 else brown)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True
title = doc.styles["Title"]
title.font.size = Pt(24)
title.font.bold = True
title.font.color.rgb = RGBColor.from_string(olive)
title.paragraph_format.space_before = Pt(6)
title.paragraph_format.space_after = Pt(8)
title.paragraph_format.keep_with_next = True
shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), cream)
title.element.get_or_add_pPr().append(shade)
subtitle = doc.styles["Subtitle"]
subtitle.font.size = Pt(12)
subtitle.font.color.rgb = RGBColor.from_string(brown)
subtitle.paragraph_format.space_after = Pt(14)
subtitle.paragraph_format.keep_with_next = True
for name, heading in (("Guidance Heading", True), ("Guidance Text", False)):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in doc.styles else doc.styles[name]
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(9.5)
    style.font.bold = heading
    style.font.color.rgb = RGBColor.from_string(olive if heading else brown)
    style.paragraph_format.space_before = Pt(4 if heading else 0)
    style.paragraph_format.space_after = Pt(2 if heading else 7)
    style.paragraph_format.keep_with_next = heading
    style.paragraph_format.left_indent = Mm(2)
    style.paragraph_format.right_indent = Mm(2)
    ppr = style.element.get_or_add_pPr()
    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), cream if heading else "F6F5EE")
    ppr.append(shade)
for name in ("Source Code", "Verbatim Char"):
    if name in doc.styles:
        doc.styles[name].font.name = "Courier New"
        doc.styles[name].font.size = Pt(9)
table = doc.styles["Table"]
pr = table.element.find(qn("w:tblPr"))
if pr is None:
    pr = OxmlElement("w:tblPr"); table.element.append(pr)
for tag in ("tblBorders", "tblCellMar"):
    old = pr.find(qn("w:" + tag))
    if old is not None: pr.remove(old)
borders = OxmlElement("w:tblBorders")
for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
    edge = OxmlElement("w:" + name)
    for key, val in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
        edge.set(qn("w:" + key), val)
    borders.append(edge)
pr.append(borders)
margins = OxmlElement("w:tblCellMar")
for name, width in (("top", "80"), ("bottom", "80"), ("left", "100"), ("right", "100")):
    edge = OxmlElement("w:" + name); edge.set(qn("w:w"), width); edge.set(qn("w:type"), "dxa")
    margins.append(edge)
pr.append(margins)
for section in doc.sections:
    section.page_width = Mm(210); section.page_height = Mm(297)
    section.top_margin = Mm(24); section.bottom_margin = Mm(22)
    section.left_margin = Mm(21); section.right_margin = Mm(21)
    section.header_distance = Mm(10); section.footer_distance = Mm(10)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.text = "Ezemvelo KZN Wildlife\tScientific Services"
    header.paragraph_format.tab_stops.add_tab_stop(Mm(168), 2)
    for run in header.runs:
        run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(brown)
    first = section.first_page_header.paragraphs[0]
    first.add_run().add_picture(str(path.parent / "style/img/corplogo120.png"), width=Mm(16))
    run = first.add_run("   EZEMVELO KZN WILDLIFE   |   Scientific Services")
    run.font.name = "Arial"; run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string(brown)
    for footer in (section.footer, section.first_page_footer):
        p = footer.paragraphs[0]
        p.text = "EKZNW | Dataset development\t"
        p.paragraph_format.tab_stops.add_tab_stop(Mm(168), 2)
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
        p._p.append(field)
        for run in p.runs:
            run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(brown)
doc.core_properties.author = "Ezemvelo KZN Wildlife"
doc.core_properties.title = "EKZNW dataset development styles"
doc.core_properties.subject = "Reusable Quarto Word reference"
doc.core_properties.comments = ""
doc.save(path)
